#!/usr/bin/env python3
"""
Phase 2 finishing pass — applied here with openpyxl instead of by hand in Excel.

What this does (everything Phase 2 can do without a live Excel calc engine):
  1. Populates the three Sensitivities grids with REAL computed values from the model
     engine (a static snapshot at the base inputs), plus a max-bid-by-target table.
  2. Adds conditional formatting so status / acceptance-check cells turn red/green.
  3. Sets print areas, fit-to-page, and freeze panes for clean printing & navigation.

What still needs real Excel (documented in PHASE2_GUIDE.md):
  - Converting the snapshot grids into LIVE two-way Data Tables (so they recompute when
    you change inputs) and running interactive Goal Seek. The numbers written here match
    what those live tables produce at the current inputs.

Computed cells are coloured PURPLE to distinguish them from inputs (blue) and formulas
(black). Run order: build_model.py -> solve_and_test.py -> phase2_finish.py.
"""

from dataclasses import replace
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.properties import PageSetupProperties

from solve_and_test import Inputs, engine, solve_max_bid, edate

WB = "MixedUse_Acquisition_Model.xlsx"

NF_PCT = '0.00%'
NF_MULT = '0.00"x"'
NF_USD = '$#,##0'
COMP = Font(color="FF7030A0")          # purple = computed snapshot
COMP_B = Font(color="FF7030A0", bold=True)
RED = PatternFill("solid", fgColor="FFFFC7CE")
GREEN = PatternFill("solid", fgColor="FFC6EFCE")
SUBFILL = PatternFill("solid", fgColor="FFE2EFDA")


def run():
    wb = load_workbook(WB)
    base = Inputs()
    se = wb["Sensitivities"]

    def put(ws, coord, val, nf=NF_PCT, font=COMP, fill=None):
        c = ws[coord]; c.value = val; c.number_format = nf; c.font = font
        if fill: c.fill = fill

    # ---------------------------------------------------------------------------------
    # Grid 1 — rows A11:A15 = exit cap ; cols B10:F10 = MF rent growth Y1 -> Levered IRR
    # ---------------------------------------------------------------------------------
    caps1 = [se[f"A{r}"].value for r in range(11, 16)]
    grs = [se[f"{c}10"].value for c in "BCDEF"]
    for r, cap in zip(range(11, 16), caps1):
        for cl, g in zip("BCDEF", grs):
            inp = replace(base, exit_cap_blended=cap,
                          vec_mfrent=(g,) + tuple(base.vec_mfrent[1:]))
            put(se, f"{cl}{r}", engine(inp, base.input_price)["irr"])

    # ---------------------------------------------------------------------------------
    # Grid 2 — rows A21:A25 = purchase price ; cols B20:F20 = exit cap -> Levered IRR
    # ---------------------------------------------------------------------------------
    prices = [se[f"A{r}"].value for r in range(21, 26)]
    caps2 = [se[f"{c}20"].value for c in "BCDEF"]
    for r, p in zip(range(21, 26), prices):
        for cl, cap in zip("BCDEF", caps2):
            put(se, f"{cl}{r}", engine(replace(base, exit_cap_blended=cap), p)["irr"])

    # ---------------------------------------------------------------------------------
    # Grid 3 — disposition date -> Levered IRR & Equity Multiple
    # ---------------------------------------------------------------------------------
    for k, r in enumerate(range(31, 36)):
        e = engine(replace(base, disp=edate(base.disp, (k - 2) * 12)), base.input_price)
        put(se, f"B{r}", e["irr"])
        put(se, f"C{r}", e["em"], nf=NF_MULT)

    # ---------------------------------------------------------------------------------
    # Max-bid by target levered IRR (the Goal-Seek deliverable, precomputed)
    # ---------------------------------------------------------------------------------
    r0 = 39
    put(se, f"A{r0}", "Max supportable bid by target levered IRR (computed snapshot):",
        nf='General', font=COMP_B)
    hdr = ["Target levered IRR", "Max bid ($)", "Going-in cap", "$/unit"]
    for j, h in enumerate(hdr):
        c = se.cell(row=r0 + 1, column=1 + j); c.value = h
        c.font = COMP_B; c.fill = SUBFILL; c.alignment = Alignment(horizontal="center", wrap_text=True)
    for k, tgt in enumerate([0.10, 0.12, 0.15, 0.18, 0.20]):
        r = r0 + 2 + k
        price = solve_max_bid(base, tgt)
        if price is None:
            put(se, f"A{r}", tgt); continue
        e = engine(base, price)
        put(se, f"A{r}", tgt, nf=NF_PCT)
        put(se, f"B{r}", round(price, 0), nf=NF_USD)
        put(se, f"C{r}", e["going_in"] / price, nf=NF_PCT)
        put(se, f"D{r}", round(price / len(base.mf_units), 0), nf=NF_USD)
    note = se.cell(row=r0 + 8, column=1)
    note.value = ("Snapshot at current base inputs. To make grids recompute live, convert "
                  "to Data Tables per PHASE2_GUIDE.md (cell inputs are noted under each grid).")
    note.font = Font(italic=True, color="FF808080", size=9)

    # ---------------------------------------------------------------------------------
    # Conditional formatting on status / acceptance-check cells
    # ---------------------------------------------------------------------------------
    def cf_bool(ws, coord):
        ws.conditional_formatting.add(coord, FormulaRule(formula=[coord], fill=GREEN))
        ws.conditional_formatting.add(coord, FormulaRule(formula=[f"NOT({coord})"], fill=RED))

    def cf_text_ok(ws, coord):
        ws.conditional_formatting.add(coord, FormulaRule(
            formula=[f'ISNUMBER(SEARCH("OK",{coord}))'], fill=GREEN))
        ws.conditional_formatting.add(coord, FormulaRule(
            formula=[f'NOT(ISNUMBER(SEARCH("OK",{coord})))'], fill=RED))

    sm = wb["Summary"]
    cf_bool(sm, "$B$3")
    for nm in ("Chk_SU", "Chk_WF", "Chk_Loan"):
        if nm in wb.defined_names:
            sh, coord = list(wb.defined_names[nm].destinations)[0]
            cf_bool(wb[sh], coord)
    if "ModelStatus" in wb.defined_names:
        sh, coord = list(wb.defined_names["ModelStatus"].destinations)[0]
        cf_text_ok(wb[sh], coord)

    # ---------------------------------------------------------------------------------
    # Print setup + freeze panes
    # ---------------------------------------------------------------------------------
    ic = wb["IC One-Pager"]
    ic.page_setup.orientation = "portrait"
    ic.page_setup.fitToWidth = 1
    ic.page_setup.fitToHeight = 1
    ic.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ic.print_options.horizontalCentered = True

    # Summary: fit to one page wide
    sm.page_setup.orientation = "portrait"
    sm.page_setup.fitToWidth = 1
    sm.page_setup.fitToHeight = 0
    sm.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sm.print_area = "A1:D60"

    # Freeze panes: keep label column(s) + header visible while scrolling the monthly grid
    monthly = ["MF Revenue", "Retail Revenue", "OpEx", "Taxes", "CapEx",
               "Debt", "Monthly CF", "Waterfall", "Assumptions"]
    for name in monthly:
        wb[name].freeze_panes = "E5"
    for name in ["Control Panel", "Summary", "Sensitivities", "IC One-Pager"]:
        wb[name].freeze_panes = "B5"

    wb.save(WB)
    print("Phase 2 finishing pass complete:")
    print("  - Sensitivity grids 1/2/3 populated with computed values")
    print("  - Max-bid-by-target table written")
    print("  - Conditional formatting on status & acceptance-check cells")
    print("  - Print areas, fit-to-page, and freeze panes set")


if __name__ == "__main__":
    run()
